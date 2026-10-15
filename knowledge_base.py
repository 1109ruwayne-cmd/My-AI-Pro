"""
knowledge_base.py - reads your files and finds the parts that answer a question.

  1. Every file in the "knowledge" folder is turned into plain text
     (PDF, Word, PowerPoint, Excel, CSV, text, Markdown, HTML, JSON ...).
  2. The text is cut into small, overlapping passages.
  3. Passages are ranked with BM25 (the classic search-engine formula).
     If all your files fit in the AI's reading budget, it simply gets everything.
     Otherwise the AI searches them itself when a question needs them
     (KnowledgeBase.tool_search), or gets the best matches for the question.

This part needs no AI service - it runs locally and instantly.
"""
from __future__ import annotations

import csv
import datetime as _dt
import html
import io
import json
import math
import re
import sys
import threading
from collections import Counter
from dataclasses import dataclass, field
from html.parser import HTMLParser
from pathlib import Path

# --------------------------------------------------------------------------- types
TEXT_TYPES = {".txt", ".md", ".markdown", ".rst", ".log", ".yaml", ".yml", ".xml", ".ini", ".cfg", ".tex"}
SUPPORTED_TYPES = TEXT_TYPES | {
    ".pdf", ".docx", ".pptx", ".xlsx", ".xlsm", ".csv", ".tsv", ".json", ".html", ".htm",
}
LEGACY_TYPES = {".doc": ".docx", ".xls": ".xlsx", ".ppt": ".pptx"}
IMAGE_TYPES = {".png", ".jpg", ".jpeg", ".gif", ".webp", ".bmp", ".tif", ".tiff", ".heic", ".svg"}

CHUNK_CHARS = 1200        # size of one passage (characters)
CHUNK_OVERLAP = 200       # neighbouring passages overlap so nothing gets cut in half
ROWS_PER_SECTION = 40     # spreadsheet rows grouped per section
MAX_ROWS_PER_SHEET = 50_000

csv.field_size_limit(min(sys.maxsize, 2**31 - 1))


@dataclass
class Section:
    location: str   # "page 3", "slide 2", "sheet 'Sales', rows 2-41" or ""
    text: str


@dataclass
class Chunk:
    file: str
    location: str
    text: str
    order: int
    tf: Counter = field(default_factory=Counter, repr=False)
    length: int = 0
    key: int = 0          # fingerprint of the text, used to skip duplicates


@dataclass
class FileInfo:
    name: str
    size: int
    mtime: float
    kind: str
    pages: int = 0
    unit: str = ""
    sections: list = field(default_factory=list, repr=False)
    chunks: list = field(default_factory=list, repr=False)
    chars: int = 0
    warning: str = ""
    error: str = ""

    def to_dict(self) -> dict:
        return {
            "name": self.name, "size": self.size, "kind": self.kind,
            "pages": self.pages, "unit": self.unit, "chunks": len(self.chunks),
            "chars": self.chars, "warning": self.warning, "error": self.error,
        }


# ------------------------------------------------------------------------ parsers
def _decode(raw: bytes) -> str:
    if raw[:2] in (b"\xff\xfe", b"\xfe\xff"):
        return raw.decode("utf-16", errors="replace")
    for enc in ("utf-8-sig", "cp1252"):
        try:
            return raw.decode(enc)
        except UnicodeDecodeError:
            continue
    return raw.decode("latin-1", errors="replace")


def _parse_text(path: Path) -> list[Section]:
    return [Section("", _decode(path.read_bytes()))]


def _parse_json(path: Path) -> list[Section]:
    raw = _decode(path.read_bytes())
    try:
        raw = json.dumps(json.loads(raw), indent=1, ensure_ascii=False)
    except ValueError:
        pass
    return [Section("", raw)]


class _HTMLText(HTMLParser):
    BLOCK = {"p", "div", "br", "li", "tr", "h1", "h2", "h3", "h4", "h5", "h6", "section",
             "article", "table", "ul", "ol", "header", "footer", "blockquote", "pre"}
    SKIP = {"script", "style", "noscript", "template", "svg"}

    def __init__(self):
        super().__init__(convert_charrefs=True)
        self.parts: list[str] = []
        self._skip = 0

    def handle_starttag(self, tag, attrs):
        if tag in self.SKIP:
            self._skip += 1
        elif tag in self.BLOCK:
            self.parts.append("\n")
        elif tag in ("td", "th"):
            self.parts.append(" | ")

    def handle_endtag(self, tag):
        if tag in self.SKIP:
            self._skip = max(0, self._skip - 1)
        elif tag in self.BLOCK:
            self.parts.append("\n")

    def handle_data(self, data):
        if not self._skip:
            self.parts.append(data)


def _parse_html(path: Path) -> list[Section]:
    parser = _HTMLText()
    parser.feed(_decode(path.read_bytes()))
    parser.close()
    return [Section("", "".join(parser.parts))]


def _parse_pdf(path: Path) -> list[Section]:
    from pypdf import PdfReader

    reader = PdfReader(str(path))
    if reader.is_encrypted:
        try:
            if not reader.decrypt(""):
                raise ValueError
        except Exception:
            raise ValueError("it is password-protected. Remove the password and add it again")
    out = []
    for i, page in enumerate(reader.pages, start=1):
        try:
            text = page.extract_text() or ""
        except Exception:
            text = ""
        out.append(Section(f"page {i}", text))
    return out


def _parse_docx(path: Path) -> list[Section]:
    import docx
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    document = docx.Document(str(path))
    parts: list[str] = []
    for child in document.element.body.iterchildren():
        tag = child.tag.rsplit("}", 1)[-1]
        if tag == "p":
            para = Paragraph(child, document)
            text = para.text.strip()
            if not text:
                continue
            try:
                style = (para.style.name or "").lower() if para.style is not None else ""
            except Exception:
                style = ""
            if style.startswith("heading"):
                digits = re.findall(r"\d+", style)
                text = "#" * min(int(digits[0]) if digits else 1, 6) + " " + text
            elif style == "title":
                text = "# " + text
            elif "list" in style:
                text = "- " + text
            parts.append(text)
        elif tag == "tbl":
            try:
                table = Table(child, document)
                rows = []
                for row in table.rows:
                    cells, prev = [], None
                    for cell in row.cells:
                        if cell._tc is prev:      # merged cells repeat - skip duplicates
                            continue
                        prev = cell._tc
                        cells.append(cell.text.strip().replace("\n", " "))
                    if any(cells):
                        rows.append(" | ".join(cells))
                if rows:
                    parts.append("\n".join(rows))
            except Exception:
                continue
    return [Section("", "\n\n".join(parts))]


def _parse_pptx(path: Path) -> list[Section]:
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    def shape_text(shape) -> list[str]:
        out: list[str] = []
        try:
            if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
                for sub in shape.shapes:
                    out.extend(shape_text(sub))
                return out
        except Exception:
            pass
        if getattr(shape, "has_text_frame", False):
            text = shape.text_frame.text.strip()
            if text:
                out.append(text)
        if getattr(shape, "has_table", False):
            for row in shape.table.rows:
                cells = [c.text.strip().replace("\n", " ") for c in row.cells]
                if any(cells):
                    out.append(" | ".join(cells))
        return out

    prs = Presentation(str(path))
    out = []
    for i, slide in enumerate(prs.slides, start=1):
        parts: list[str] = []
        for shape in slide.shapes:
            parts.extend(shape_text(shape))
        try:
            if slide.has_notes_slide:
                notes = slide.notes_slide.notes_text_frame.text.strip()
                if notes:
                    parts.append("Speaker notes: " + notes)
        except Exception:
            pass
        out.append(Section(f"slide {i}", "\n".join(parts)))
    return out


def _fmt(value) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    if isinstance(value, _dt.datetime):
        if (value.hour, value.minute, value.second) == (0, 0, 0):
            return value.date().isoformat()
        return value.isoformat(sep=" ", timespec="minutes")
    if isinstance(value, (_dt.date, _dt.time)):
        return value.isoformat()
    return str(value).strip()


def _table_sections(rows: list[tuple[int, list[str]]], label: str) -> list[Section]:
    """Turn spreadsheet rows into readable text like 'Row 5: Name: Thabo; Amount: 500'."""
    if not rows:
        return []
    header_row, header = rows[0]
    header = [h.strip() for h in header]
    named = sum(1 for h in header if h) >= max(1, len(header) // 2)
    body = rows[1:] if named else rows
    prefix = label + ", " if label else ""

    def col(k: int) -> str:
        return header[k] if k < len(header) and header[k] else "column %d" % (k + 1)

    if not body:
        return [Section("%srow %d" % (prefix, header_row), " | ".join(header))]
    out = []
    for start in range(0, len(body), ROWS_PER_SECTION):
        group = body[start:start + ROWS_PER_SECTION]
        lines = []
        if named:
            lines.append("Columns: " + " | ".join(col(k) for k in range(len(header))))
        for r_idx, vals in group:
            if named:
                pairs = ["%s: %s" % (col(k), v) for k, v in enumerate(vals) if v.strip()]
                lines.append("Row %d: %s" % (r_idx, "; ".join(pairs)))
            else:
                lines.append("Row %d: %s" % (r_idx, " | ".join(vals)))
        out.append(Section("%srows %d-%d" % (prefix, group[0][0], group[-1][0]), "\n".join(lines)))
    return out


def _parse_xlsx(path: Path) -> list[Section]:
    from openpyxl import load_workbook

    wb = load_workbook(str(path), read_only=True, data_only=True)
    out: list[Section] = []
    try:
        for ws in wb.worksheets:
            rows = []
            for r_idx, row in enumerate(ws.iter_rows(values_only=True), start=1):
                if r_idx > MAX_ROWS_PER_SHEET:
                    break
                vals = [_fmt(v) for v in row]
                while vals and not vals[-1]:
                    vals.pop()
                if any(vals):
                    rows.append((r_idx, vals))
            out.extend(_table_sections(rows, "sheet '%s'" % ws.title))
    finally:
        wb.close()
    return out


def _parse_csv(path: Path, delimiter: str | None = None) -> list[Section]:
    text = _decode(path.read_bytes())
    if delimiter is None:
        try:
            delimiter = csv.Sniffer().sniff(text[:20000], delimiters=",;\t|").delimiter
        except csv.Error:
            delimiter = ","
    rows = []
    for r_idx, row in enumerate(csv.reader(io.StringIO(text), delimiter=delimiter), start=1):
        if r_idx > MAX_ROWS_PER_SHEET:
            break
        vals = [c.strip() for c in row]
        while vals and not vals[-1]:
            vals.pop()
        if any(vals):
            rows.append((r_idx, vals))
    return _table_sections(rows, "")


PARSERS = {
    ".pdf": _parse_pdf, ".docx": _parse_docx, ".pptx": _parse_pptx,
    ".xlsx": _parse_xlsx, ".xlsm": _parse_xlsx,
    ".csv": _parse_csv, ".tsv": lambda p: _parse_csv(p, "\t"),
    ".json": _parse_json, ".html": _parse_html, ".htm": _parse_html,
}
for _ext in TEXT_TYPES:
    PARSERS[_ext] = _parse_text

UNITS = {".pdf": "page", ".pptx": "slide"}


# ----------------------------------------------------------------- text helpers
def clean_text(text: str) -> str:
    text = text.replace("\r\n", "\n").replace("\r", "\n").replace("\x00", "")
    text = re.sub(r"([a-z])-\n([a-z])", r"\1\2", text)      # words hyphenated across lines
    text = re.sub(r"[ \t\f\v\u00a0]+", " ", text)
    text = re.sub(r" ?\n ?", "\n", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def split_text(text: str, size: int = CHUNK_CHARS, overlap: int = CHUNK_OVERLAP) -> list[str]:
    """Cut text into passages of about `size` characters, breaking at natural points."""
    if len(text) <= size:
        return [text] if text.strip() else []
    chunks, start, n = [], 0, len(text)
    while start < n:
        end = min(start + size, n)
        if end < n:
            window = text[start:end]
            for sep in ("\n\n", "\n", ". ", "? ", "! ", "; ", ", ", " "):
                idx = window.rfind(sep)
                if idx > size * 0.5:
                    end = start + idx + len(sep)
                    break
        piece = text[start:end].strip()
        if piece:
            chunks.append(piece)
        if end >= n:
            break
        nxt = max(end - overlap, start + 1)
        space = text.find(" ", nxt, end)
        start = space + 1 if space != -1 else nxt
    return chunks


_TABLE_LINE = re.compile(r"^\s*\|")
_SEPARATOR_LINE = re.compile(r"^\s*\|?\s*:?-{3,}:?\s*(\|\s*:?-{3,}:?\s*)*\|?\s*$")
_ROW_LINE = re.compile(r"^Row \d+:")


def looks_like_records(text: str) -> bool:
    """True for big tables/lists where every line is its own record (e.g. an API directory)."""
    lines = [line for line in text.split("\n") if line.strip()]
    if len(lines) < 20:
        return False
    records = sum(1 for line in lines if _TABLE_LINE.match(line) or _ROW_LINE.match(line))
    return records / len(lines) >= 0.6


def split_records(text: str, short: int = 100, target: int = 300) -> list[str]:
    """One passage per table row, so each record can be found on its own.
    Only very short rows (under `short` characters) are grouped together."""
    chunks, buf = [], ""
    for line in text.split("\n"):
        line = line.strip()
        if not line or _SEPARATOR_LINE.match(line):
            continue
        if len(line) >= short:
            if buf:
                chunks.append(buf)
                buf = ""
            chunks.extend(split_text(line) if len(line) > CHUNK_CHARS else [line])
        elif buf and len(buf) + 1 + len(line) > target:
            chunks.append(buf)
            buf = line
        else:
            buf = buf + "\n" + line if buf else line
    if buf:
        chunks.append(buf)
    return chunks


_FIRST_LINK = re.compile(r"\]\((https?://[^\s)]+)\)")


def fingerprint(text: str, record: bool = False) -> int:
    """Same fingerprint = same passage. A single directory row is identified by its link,
    so one API listed in several categories is only shown once."""
    if record and "\n" not in text:
        m = _FIRST_LINK.search(text)
        if m:
            return hash("link:" + m.group(1).lower())
    return hash(" ".join(text.lower().split()))


def pretty(text: str) -> str:
    """Make table rows readable: '| [Name](link) | What it does |' -> '[Name](link) — What it does'."""
    out = []
    for line in text.split("\n"):
        if _TABLE_LINE.match(line):
            line = re.sub(r"^\s*\|\s*|\s*\|\s*$", "", line).replace(" | ", " — ")
        out.append(html.unescape(line))
    return "\n".join(out)


def short_label(text: str) -> str:
    """A short name for a passage: the first link text (e.g. an API name) or its first words."""
    m = re.search(r"\[([^\]\n]{2,80})\]\(", text)
    if m:
        return html.unescape(m.group(1)).strip()
    first = " ".join(re.sub(r"[|#*_`>\[\]]", " ", text.split("\n", 1)[0]).split())
    return first[:48] + "…" if len(first) > 48 else first


STOPWORDS = frozenset("""
a about above after again against all am an and any are as at be because been before being below
between both but by can could did do does doing down during each few for from further had has have
having he her here hers herself him himself his how i if in into is it its itself just me more most
my myself no nor not now of off on once only or other our ours ourselves out over own same she should
so some such than that the their theirs them themselves then there these they this those through to
too under until up very was we were what when where which while who whom why will with would you
your yours yourself yourselves please tell give show find list explain describe according file files
document documents doc
""".split())

_WORD = re.compile(r"\w+", re.UNICODE)


def _stem(word: str) -> str:
    if word.isdigit() or len(word) <= 3:
        return word
    if word.endswith("ies") and len(word) > 4:
        return word[:-3] + "y"
    if word.endswith("ing") and len(word) > 5:
        return word[:-3]
    if word.endswith("ed") and len(word) > 4:
        return word[:-2]
    if word.endswith("es") and len(word) > 4 and word[-3] in "sxz":
        return word[:-2]
    if word.endswith("s") and not word.endswith(("ss", "us", "is")):
        return word[:-1]
    return word


def tokenize(text: str) -> list[str]:
    out = []
    for word in _WORD.findall(text.lower()):
        if word in STOPWORDS or (len(word) < 2 and not word.isdigit()):
            continue
        out.append(sys.intern(_stem(word)))  # intern: big knowledge bases use far less memory
    return out


def make_snippet(text: str, terms: list[str], width: int = 240) -> str:
    """Pick the part of a passage that contains the most search words."""
    low = text.lower()
    best_start, best_hits = 0, -1
    candidates = {0}
    for t in set(terms):
        pos, found = low.find(t), 0
        while pos != -1 and found < 3:
            candidates.add(max(0, pos - 60))
            pos, found = low.find(t, pos + 1), found + 1
    for start in sorted(candidates):
        window = low[start:start + width]
        hits = sum(1 for t in set(terms) if t in window)
        if hits > best_hits:
            best_start, best_hits = start, hits
    snippet = text[best_start:best_start + width].replace("\n", " ").strip()
    if best_start > 0:
        snippet = "…" + snippet
    if best_start + width < len(text):
        snippet += "…"
    return snippet


# ------------------------------------------------------------------- one file
def load_file(path: Path, size: int, mtime: float) -> FileInfo:
    ext = path.suffix.lower()
    info = FileInfo(name=path.name, size=size, mtime=mtime, kind=ext.lstrip(".") or "file",
                    unit=UNITS.get(ext, ""))
    if ext in LEGACY_TYPES:
        info.warning = f"Old {ext} format. Open it and use 'Save as' {LEGACY_TYPES[ext]}, then add it again."
        return info
    if ext in IMAGE_TYPES:
        info.warning = "Pictures can't be read yet. Only files that contain text."
        return info
    if ext not in PARSERS:
        info.warning = "This file type isn't supported."
        return info
    try:
        sections = PARSERS[ext](path)
    except Exception as exc:  # broken / protected files must not crash the app
        info.error = f"Couldn't read this file: {exc}"
        return info

    title_terms = tokenize(path.stem.replace("_", " ").replace("-", " "))
    order = 0
    for sec in sections:
        text = clean_text(sec.text)
        if not text:
            continue
        info.sections.append(Section(sec.location, text))
        info.chars += len(text)
        record = looks_like_records(text)
        pieces = split_records(text) if record else split_text(text)
        for piece in pieces:
            terms = tokenize(piece) + title_terms
            info.chunks.append(Chunk(info.name, sec.location, piece, order, Counter(terms), len(terms),
                                     fingerprint(piece, record)))
            order += 1

    if ext in (".pdf", ".pptx"):
        info.pages = len(sections)
    elif ext in (".xlsx", ".xlsm"):
        info.pages = len({s.location.split(",")[0] for s in sections})
        info.unit = "sheet"

    if info.chars == 0:
        info.warning = "No readable text found." + (
            " It looks like a scanned PDF (pictures of pages)." if ext == ".pdf" else "")
    elif ext == ".pdf" and info.chars / max(1, info.pages) < 40:
        info.warning = "Very little text found. Some pages may be scanned pictures."
    return info


# ------------------------------------------------------------- knowledge base
class KnowledgeBase:
    def __init__(self, folder: Path):
        self.folder = Path(folder)
        self.folder.mkdir(parents=True, exist_ok=True)
        self.lock = threading.RLock()
        self.files: dict[str, FileInfo] = {}
        self.chunks: list[Chunk] = []
        self.df: Counter = Counter()
        self.avgdl = 0.0
        self.total_chars = 0
        self._signature = None
        self.refresh(force=True)

    # Re-read the folder only when something changed (new, edited or deleted files).
    def _scan(self) -> tuple:
        entries = []
        for p in sorted(self.folder.iterdir(), key=lambda p: p.name.lower()):
            if p.is_file() and not p.name.startswith((".", "~$")):
                st = p.stat()
                entries.append((p.name, st.st_size, st.st_mtime))
        return tuple(entries)

    def refresh(self, force: bool = False) -> bool:
        with self.lock:
            sig = self._scan()
            if not force and sig == self._signature:
                return False
            files = {}
            for name, size, mtime in sig:
                old = self.files.get(name)
                if old and old.size == size and old.mtime == mtime:
                    files[name] = old
                else:
                    files[name] = load_file(self.folder / name, size, mtime)
            chunks = [c for f in files.values() for c in f.chunks]
            df: Counter = Counter()
            total_len = 0
            for c in chunks:
                df.update(c.tf.keys())
                total_len += c.length
            self.files, self.chunks, self.df = files, chunks, df
            self.avgdl = total_len / len(chunks) if chunks else 0.0
            self.total_chars = sum(f.chars for f in files.values())
            self._signature = sig
            return True

    def list_files(self) -> list[dict]:
        with self.lock:
            return [f.to_dict() for f in self.files.values()]

    def search(self, query: str) -> list[tuple[float, Chunk]]:
        """BM25 ranking of all passages for the query (best first)."""
        with self.lock:
            q = Counter(tokenize(query))
            if not q or not self.chunks:
                return []
            n, k1, b = len(self.chunks), 1.5, 0.75
            idf = {}
            for t in q:
                d = self.df.get(t, 0)
                if d:
                    idf[t] = math.log(1 + (n - d + 0.5) / (d + 0.5))
            if not idf:
                return []
            avg = self.avgdl or 1.0
            scored = []
            for c in self.chunks:
                s = 0.0
                for t, w in idf.items():
                    tf = c.tf.get(t)
                    if tf:
                        s += w * min(q[t], 3) * tf * (k1 + 1) / (tf + k1 * (1 - b + b * c.length / avg))
                if s > 0:
                    scored.append((s, c))
            scored.sort(key=lambda x: -x[0])
            return scored

    def build_context(self, query: str, budget_chars: int) -> dict:
        """Choose what the AI reads for this question."""
        with self.lock:
            files = list(self.files.values())
            file_notes = [(f.name, "" if f.chunks else "no readable text") for f in files]
            ranked = self.search(query)
            terms = list(set(tokenize(query)))
            sources = self._sources(ranked, terms)
            if not self.chunks:
                return {"mode": "empty", "blocks": [], "sources": [], "files": file_notes}

            if self.total_chars <= budget_chars:
                # Everything fits: give the AI every file in full.
                blocks = [(f.name, s.location, s.text) for f in files for s in f.sections]
                return {"mode": "all", "blocks": blocks, "sources": sources, "files": file_notes}

            # Too much text: best-matching passages first (identical passages only once)...
            picked, seen, used = [], set(), 0
            for _, c in ranked:
                if c.key in seen or used + len(c.text) > budget_chars:
                    continue
                picked.append(c)
                seen.add(c.key)
                used += len(c.text)
                if used > budget_chars * 0.97:
                    break
            # ...and for broad questions ("summarise everything") add the start of every file.
            if used < budget_chars * 0.6:
                iters = [iter(f.chunks) for f in files if f.chunks]
                while iters and used < budget_chars * 0.95:
                    for it in list(iters):
                        c = next(it, None)
                        if c is None or used + len(c.text) > budget_chars:
                            iters.remove(it)
                            continue
                        if c.key in seen:
                            continue
                        picked.append(c)
                        seen.add(c.key)
                        used += len(c.text)
            picked.sort(key=lambda c: (c.file.lower(), c.order))
            blocks = [(c.file, c.location, c.text) for c in picked]
            return {"mode": "search", "blocks": blocks, "sources": sources, "files": file_notes}

    def tool_search(self, query: str, max_results: int = 20, budget_chars: int = 15000) -> tuple[str, list[dict]]:
        """Search on behalf of the AI (its `search_files` tool).
        Returns (text the AI reads, source cards for the chat window)."""
        with self.lock:
            ranked = self.search(query)
            terms = list(set(tokenize(query)))
            picked, seen, used = [], set(), 0
            for _, c in ranked:
                if c.key in seen:
                    continue
                text = pretty(c.text)
                if picked and used + len(text) > budget_chars:
                    break
                picked.append((c, text))
                seen.add(c.key)
                used += len(text)
                if len(picked) >= max_results:
                    break
            sources = self._sources(ranked, terms)
        if not picked:
            return (f'No results for "{query}". Try fewer, different or more general keywords '
                    "(synonyms, singular words, the English term).", [])
        more = "" if len(picked) < max_results else " There may be more: search again with other keywords or ask for more results."
        lines = [f'Results for "{query}" (best match first, {len(picked)} shown).{more}']
        for i, (c, text) in enumerate(picked, 1):
            where = c.file + (f", {c.location}" if c.location else "")
            lines.append(f"[{i}] ({where})\n{text}")
        return "\n\n".join(lines), sources

    @staticmethod
    def _sources(ranked: list, terms: list[str], limit: int = 6) -> list[dict]:
        if not ranked:
            return []
        top = ranked[0][0]
        out, seen = [], set()
        for score, c in ranked:
            if score < top * 0.3 or len(out) >= limit:
                break
            # one source per page/slide; for files without pages, one per distinct passage
            place = (c.file, c.location) if c.location else None
            if c.key in seen or (place and place in seen):
                continue
            seen.add(c.key)
            if place:
                seen.add(place)
            out.append({
                "file": c.file, "location": c.location,
                "label": "" if c.location else short_label(c.text),
                "snippet": make_snippet(pretty(c.text), terms), "passage": pretty(c.text)[:1500],
                "score": round(score, 2), "key": str(c.key),
            })
        return out
